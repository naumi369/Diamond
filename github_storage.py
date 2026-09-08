"""
Persist the master diamond inventory on GitHub.

Required Streamlit secrets (Settings → Secrets on share.streamlit.io):

    GITHUB_TOKEN = "ghp_...."          # classic PAT with `repo` scope
    GITHUB_REPO  = "username/diamond"  # owner/repo
    GITHUB_BRANCH = "main"             # optional, default main
    GITHUB_PATH  = "data/master_diamonds.xlsx"  # optional

Or a single TOML block:

    [github]
    token = "ghp_...."
    repo = "username/diamond"
    branch = "main"
    path = "data/master_diamonds.xlsx"
"""

from __future__ import annotations

import base64
import io
from datetime import datetime, timezone
from typing import Any, Dict, Optional, Tuple

import pandas as pd
import requests

API = "https://api.github.com"


def _secrets() -> Dict[str, str]:
    """Read GitHub settings from st.secrets (works only inside Streamlit)."""
    try:
        import streamlit as st
        sec = st.secrets
    except Exception:
        return {}

    # Flat keys
    token = sec.get("GITHUB_TOKEN") or sec.get("github_token")
    repo = sec.get("GITHUB_REPO") or sec.get("github_repo")
    branch = sec.get("GITHUB_BRANCH") or sec.get("github_branch") or "main"
    path = sec.get("GITHUB_PATH") or sec.get("github_path") or "data/master_diamonds.xlsx"
    prefs_path = (
        sec.get("GITHUB_PREFS_PATH")
        or sec.get("github_prefs_path")
        or "data/ui_prefs.json"
    )

    # Nested [github] table
    if not token and "github" in sec:
        g = sec["github"]
        token = g.get("token") or token
        repo = g.get("repo") or repo
        branch = g.get("branch") or branch
        path = g.get("path") or path
        prefs_path = g.get("prefs_path") or prefs_path

    out = {}
    if token:
        out["token"] = str(token).strip()
    if repo:
        out["repo"] = str(repo).strip()
    if branch:
        out["branch"] = str(branch).strip()
    if path:
        out["path"] = str(path).strip()
    if prefs_path:
        out["prefs_path"] = str(prefs_path).strip()
    return out


def is_configured() -> bool:
    s = _secrets()
    return bool(s.get("token") and s.get("repo"))


def _headers(token: str) -> Dict[str, str]:
    return {
        "Authorization": f"Bearer {token}",
        "Accept": "application/vnd.github+json",
        "X-GitHub-Api-Version": "2022-11-28",
        "User-Agent": "DiamondInventoryCompiler/1.0",
    }


def _get_file(repo: str, path: str, branch: str, token: str) -> Tuple[Optional[bytes], Optional[str]]:
    """
    Return (content_bytes, sha) for an existing file, or (None, None) if missing.
    """
    url = f"{API}/repos/{repo}/contents/{path}"
    r = requests.get(url, headers=_headers(token), params={"ref": branch}, timeout=60)
    if r.status_code == 404:
        return None, None
    r.raise_for_status()
    data = r.json()
    sha = data.get("sha")
    content_b64 = data.get("content", "")
    # GitHub returns base64 with newlines
    raw = base64.b64decode(content_b64.replace("\n", ""))
    return raw, sha


def load_master_excel() -> Tuple[Optional[pd.DataFrame], str]:
    """
    Load master inventory from GitHub.
    Returns (dataframe_or_None, status_message).
    """
    s = _secrets()
    if not s.get("token") or not s.get("repo"):
        return None, "GitHub secrets not configured"

    try:
        raw, sha = _get_file(s["repo"], s["path"], s["branch"], s["token"])
        if raw is None:
            return None, f"No file yet at {s['path']} on {s['branch']}"
        df = pd.read_excel(io.BytesIO(raw))
        return df, f"Loaded {len(df)} stones from GitHub ({s['path']})"
    except Exception as e:
        return None, f"GitHub load failed: {e}"


def save_master_excel(df: pd.DataFrame) -> str:
    """
    Commit the dataframe as Excel to GitHub (create or update).
    Returns a human-readable status message.
    """
    s = _secrets()
    if not s.get("token") or not s.get("repo"):
        return "GitHub secrets not configured – cannot save to cloud"

    if df is None or df.empty:
        return "Nothing to save (inventory is empty)"

    try:
        buf = io.BytesIO()
        out = df.copy()
        # Strip trailing .0 and keep as plain text IDs
        import re as _re
        def _id_str(v):
            if v is None:
                return None
            try:
                import pandas as _pd
                if _pd.isna(v):
                    return None
            except Exception:
                pass
            if isinstance(v, float) and v == int(v):
                return str(int(v))
            if isinstance(v, int) and not isinstance(v, bool):
                return str(v)
            s = str(v).strip()
            if s.lower() in ("", "nan", "none", "<na>"):
                return None
            if _re.match(r"^-?\d+\.0+$", s):
                s = s.split(".")[0]
            return s
        for col in ("certificate_no", "stock_no"):
            if col in out.columns:
                out[col] = out[col].map(_id_str)
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            out.to_excel(writer, index=False, sheet_name="Inventory")
            # Force text format so Excel does not turn IDs into numbers (… .0)
            ws = writer.sheets["Inventory"]
            headers = {cell.value: cell.column for cell in ws[1]}
            from openpyxl.styles import numbers
            for col_name in ("certificate_no", "stock_no"):
                if col_name not in headers:
                    continue
                col_idx = headers[col_name]
                for row in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row, column=col_idx)
                    if cell.value is not None:
                        cell.number_format = numbers.FORMAT_TEXT
                        cell.value = str(cell.value)
        content_b64 = base64.b64encode(buf.getvalue()).decode("ascii")

        # Need current SHA to update an existing file
        _, sha = _get_file(s["repo"], s["path"], s["branch"], s["token"])

        ts = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
        message = f"Update master diamond inventory ({len(df)} stones) – {ts}"

        payload: Dict[str, Any] = {
            "message": message,
            "content": content_b64,
            "branch": s["branch"],
        }
        if sha:
            payload["sha"] = sha

        url = f"{API}/repos/{s['repo']}/contents/{s['path']}"
        r = requests.put(url, headers=_headers(s["token"]), json=payload, timeout=90)
        if r.status_code in (200, 201):
            commit = r.json().get("commit", {}).get("html_url", "")
            return f"Saved {len(df)} stones to GitHub → {s['path']}" + (
                f"\n{commit}" if commit else ""
            )
        return f"GitHub save failed ({r.status_code}): {r.text[:500]}"
    except Exception as e:
        return f"GitHub save failed: {e}"


# ---------------------------------------------------------------------------
# UI preferences (column order, etc.)
# ---------------------------------------------------------------------------

import json


def load_ui_prefs() -> Dict[str, Any]:
    """Load UI prefs JSON from GitHub. Returns {} if missing/unavailable."""
    s = _secrets()
    if not s.get("token") or not s.get("repo"):
        return {}
    try:
        raw, _sha = _get_file(s["repo"], s["prefs_path"], s["branch"], s["token"])
        if raw is None:
            return {}
        return json.loads(raw.decode("utf-8"))
    except Exception:
        return {}


def save_ui_prefs(prefs: Dict[str, Any]) -> str:
    """Save UI prefs JSON to GitHub."""
    s = _secrets()
    if not s.get("token") or not s.get("repo"):
        return "GitHub secrets not configured – cannot save preferences"

    try:
        content_b64 = base64.b64encode(
            json.dumps(prefs, indent=2).encode("utf-8")
        ).decode("ascii")
        _, sha = _get_file(s["repo"], s["prefs_path"], s["branch"], s["token"])
        ts = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
        payload: Dict[str, Any] = {
            "message": f"Update UI preferences – {ts}",
            "content": content_b64,
            "branch": s["branch"],
        }
        if sha:
            payload["sha"] = sha
        url = f"{API}/repos/{s['repo']}/contents/{s['prefs_path']}"
        r = requests.put(url, headers=_headers(s["token"]), json=payload, timeout=60)
        if r.status_code in (200, 201):
            return f"Preferences saved to GitHub → {s['prefs_path']}"
        return f"Prefs save failed ({r.status_code}): {r.text[:300]}"
    except Exception as e:
        return f"Prefs save failed: {e}"
