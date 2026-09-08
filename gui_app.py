#!/usr/bin/env python3
"""
Diamond Inventory Compiler – Streamlit GUI
------------------------------------------
Designed for local use AND Streamlit Community Cloud deployment.

Run locally:
    streamlit run gui_app.py

Deploy: push this folder to GitHub and connect it at share.streamlit.io
"""

from __future__ import annotations

import hashlib
import io
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import pandas as pd
import streamlit as st

# ---------------------------------------------------------------------------
# Import core compiler (same folder)
# ---------------------------------------------------------------------------
sys.path.insert(0, str(Path(__file__).resolve().parent))
from diamond_compiler import (  # noqa: E402
    STANDARD_COLUMNS,
    choose_parser,
    clean_value,
    download_videos,
    is_url,
)
from github_storage import (  # noqa: E402
    is_configured as github_configured,
    load_master_excel,
    save_master_excel,
)

# ---------------------------------------------------------------------------
# Page setup
# ---------------------------------------------------------------------------
st.set_page_config(
    page_title="Diamond Inventory Compiler",
    page_icon="💎",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _empty_inventory() -> pd.DataFrame:
    return pd.DataFrame(columns=STANDARD_COLUMNS)


def _normalize_df(df: pd.DataFrame) -> pd.DataFrame:
    """Ensure all standard columns exist and types are sane."""
    for col in STANDARD_COLUMNS:
        if col not in df.columns:
            df[col] = None
    df = df[STANDARD_COLUMNS].copy()
    for col in ("certificate_no", "stock_no"):
        df[col] = df[col].map(clean_value)
        # pandas StringDtype, blanks as <NA>
        df[col] = df[col].astype("string")
    return df


def _dedup(df: pd.DataFrame) -> pd.DataFrame:
    """Prefer certificate_no, then stock_no, keep first occurrence."""
    if df.empty:
        return df
    keys = []
    for _, row in df.iterrows():
        k = row.get("certificate_no")
        if pd.isna(k) or str(k).strip() in ("", "nan", "<NA>"):
            k = row.get("stock_no")
        if pd.isna(k) or str(k).strip() in ("", "nan", "<NA>"):
            k = f"_row_{hashlib.md5(str(row.values).encode()).hexdigest()[:12]}"
        keys.append(str(k).strip())
    df = df.copy()
    df["_dedup_key"] = keys
    df = df.drop_duplicates(subset=["_dedup_key"], keep="first")
    return df.drop(columns=["_dedup_key"])


def _parse_uploaded_excels(files):
    """Parse a list of Streamlit UploadedFile objects."""
    records = []
    logs = []
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        for uf in files:
            local = tmp_path / uf.name
            local.write_bytes(uf.getvalue())
            try:
                recs = choose_parser(local)
                logs.append(f"✓ {uf.name} → {len(recs)} stones")
                records.extend(recs)
            except Exception as e:
                logs.append(f"✗ {uf.name} failed: {e}")
    return records, logs


def _df_to_excel_bytes(df: pd.DataFrame) -> bytes:
    buf = io.BytesIO()
    out = df.copy()
    for col in ("certificate_no", "stock_no"):
        if col in out.columns:
            out[col] = out[col].map(clean_value)
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        out.to_excel(writer, index=False, sheet_name="Inventory")
        ws = writer.sheets["Inventory"]
        headers = {cell.value: cell.column for cell in ws[1]}
        from openpyxl.styles import numbers
        for col_name in ("certificate_no", "stock_no"):
            if col_name not in headers:
                continue
            col_idx = headers[col_name]
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=col_idx)
                if cell.value is not None:
                    cell.number_format = numbers.FORMAT_TEXT
                    cell.value = str(cell.value)
    return buf.getvalue()


def _df_to_csv_bytes(df: pd.DataFrame) -> bytes:
    return df.to_csv(index=False).encode("utf-8")


# ---------------------------------------------------------------------------
# Session state bootstrap
# ---------------------------------------------------------------------------
if "inventory" not in st.session_state:
    st.session_state.inventory = _empty_inventory()
if "last_logs" not in st.session_state:
    st.session_state.last_logs = []
if "video_dir" not in st.session_state:
    st.session_state.video_dir = str(
        Path(__file__).resolve().parent / "output" / "videos"
    )
if "github_loaded" not in st.session_state:
    st.session_state.github_loaded = False
if "github_status" not in st.session_state:
    st.session_state.github_status = ""

# Auto-load master inventory from GitHub once when secrets are present
if not st.session_state.github_loaded and github_configured():
    df_gh, msg = load_master_excel()
    st.session_state.github_status = msg
    st.session_state.github_loaded = True
    if df_gh is not None and not df_gh.empty:
        st.session_state.inventory = _normalize_df(df_gh)

# ---------------------------------------------------------------------------
# Sidebar
# ---------------------------------------------------------------------------
with st.sidebar:
    st.title("💎 Diamond Compiler")
    st.markdown("---")

    # GitHub cloud status
    if github_configured():
        st.success("☁️ GitHub storage connected")
        if st.session_state.github_status:
            st.caption(st.session_state.github_status)
        n = len(st.session_state.inventory)
        st.metric("Stones in inventory", n)
        if st.button("💾 Save inventory to GitHub", type="primary", use_container_width=True):
            with st.spinner("Saving to GitHub…"):
                msg = save_master_excel(st.session_state.inventory)
            st.session_state.github_status = msg
            st.info(msg)
        if st.button("🔄 Reload from GitHub", use_container_width=True):
            df_gh, msg = load_master_excel()
            st.session_state.github_status = msg
            if df_gh is not None and not df_gh.empty:
                st.session_state.inventory = _normalize_df(df_gh)
                st.success(msg)
            else:
                st.warning(msg)
            st.rerun()
    else:
        st.warning("☁️ GitHub not configured")
        with st.expander("How to save inventory on Cloud", expanded=True):
            st.markdown(
                """
Add these **Secrets** in Streamlit Cloud  
(App → Settings → Secrets):

```toml
GITHUB_TOKEN = "ghp_your_token_here"
GITHUB_REPO  = "youruser/your-repo"
GITHUB_BRANCH = "main"
GITHUB_PATH  = "data/master_diamonds.xlsx"
```

Create a GitHub **Personal Access Token** (classic) with the `repo` scope.
Then the app can load & save the master Excel in your repo automatically.
                """
            )

    st.markdown("---")
    st.markdown(
        "**Workflow**\n\n"
        "1. Upload new customer sheets\n"
        "2. Click **Process & Add**\n"
        "3. Click **Save inventory to GitHub**\n"
        "4. Next visit: inventory loads from GitHub automatically"
    )
    st.markdown("---")
    with st.expander("Video files note"):
        st.markdown(
            "Video **files** cannot stay on Streamlit Cloud (disk is temporary). "
            "Video **links** are stored in the inventory. "
            "Download files only when running locally."
        )
    st.markdown("---")
    if st.button("🗑️ Clear inventory (session only)", use_container_width=True):
        st.session_state.inventory = _empty_inventory()
        st.session_state.last_logs = []
        st.rerun()

# ---------------------------------------------------------------------------
# Tabs
# ---------------------------------------------------------------------------
tab_inventory, tab_upload, tab_summary = st.tabs(
    ["📋 Full Inventory", "📤 Upload & Add", "📊 Summary"]
)

# ========================= TAB: Upload & Add =============================
with tab_upload:
    st.header("Upload sheets → add to inventory")

    st.subheader("New customer / supplier sheets")
    st.caption("Upload one or more Excel lists from customers or suppliers.")
    new_files = st.file_uploader(
        "Upload diamond Excel files",
        type=["xlsx", "xls"],
        accept_multiple_files=True,
        key="new_uploader",
    )

    col_a, col_b = st.columns(2)
    with col_a:
        do_videos = st.checkbox(
            "Download videos (local only)",
            value=False,
            help=(
                "Works on your own computer. On Streamlit Cloud the files "
                "are temporary and disappear when the app restarts."
            ),
        )
    with col_b:
        replace_mode = st.checkbox(
            "Replace inventory instead of adding",
            value=False,
            help="If checked, the current inventory is wiped before adding these files.",
        )

    with st.expander("Optional: load previous Master Excel", expanded=False):
        st.caption(
            "Only needed if GitHub storage is not set up. "
            "Upload a previously downloaded master file to merge into."
        )
        master_file = st.file_uploader(
            "Master inventory Excel (optional)",
            type=["xlsx", "xls", "csv"],
            key="master_uploader",
        )

    if st.button("🚀 Process & Add to Inventory", type="primary", use_container_width=True):
        if master_file is None and not new_files:
            st.warning("Please upload at least one file.")
            st.stop()

        progress = st.progress(0, text="Working…")
        all_new = []
        logs = []

        # --- Load master if provided ---
        master_df = _empty_inventory()
        if master_file is not None:
            progress.progress(0.1, text="Loading master inventory…")
            try:
                if master_file.name.lower().endswith(".csv"):
                    master_df = pd.read_csv(master_file)
                else:
                    master_df = pd.read_excel(master_file)
                master_df = _normalize_df(master_df)
                logs.append(f"✓ Master loaded → {len(master_df)} stones")
            except Exception as e:
                logs.append(f"✗ Master failed: {e}")
                st.error(f"Could not read master file: {e}")
                st.stop()

        # --- Parse new sheets ---
        if new_files:
            progress.progress(0.3, text="Parsing new sheets…")
            recs, parse_logs = _parse_uploaded_excels(new_files)
            logs.extend(parse_logs)
            all_new.extend(recs)

        progress.progress(0.6, text="Merging & deduplicating…")
        new_df = (
            pd.DataFrame(all_new, columns=STANDARD_COLUMNS)
            if all_new
            else _empty_inventory()
        )
        new_df = _normalize_df(new_df)

        if replace_mode:
            combined = new_df if not new_df.empty else master_df
        else:
            pieces = []
            if not master_df.empty:
                pieces.append(master_df)
            if not st.session_state.inventory.empty and master_file is None:
                pieces.append(st.session_state.inventory)
            if not new_df.empty:
                pieces.append(new_df)
            combined = (
                pd.concat(pieces, ignore_index=True) if pieces else _empty_inventory()
            )

        before = len(combined)
        combined = _dedup(combined)
        after = len(combined)
        logs.append(
            f"Merged total: {after} unique stones (removed {before - after} duplicates)"
        )

        st.session_state.inventory = combined
        st.session_state.last_logs = logs

        # --- Optional video download (local) ---
        if do_videos and not combined.empty:
            progress.progress(0.8, text="Downloading videos…")
            video_dir = Path(st.session_state.video_dir)
            video_dir.mkdir(parents=True, exist_ok=True)
            records = combined.to_dict(orient="records")
            with_links = [r for r in records if is_url(r.get("video_link"))]
            if with_links:
                vbar = st.progress(0, text="Videos…")

                def vprog(cur, tot, name):
                    vbar.progress(cur / max(tot, 1), text=f"Video {cur}/{tot}: {name}")

                stats = download_videos(with_links, video_dir, progress_callback=vprog)
                logs.append(
                    f"Videos → {stats.get('downloaded', 0)} new, "
                    f"{stats.get('already', 0)} already, "
                    f"{stats.get('failed', 0)} failed"
                )
                vbar.empty()
            else:
                logs.append("No video URLs found to download.")

        progress.progress(1.0, text="Done!")
        st.session_state.last_logs = logs
        st.success(f"Inventory now has **{len(st.session_state.inventory)}** unique stones.")
        for line in logs:
            st.text(line)

        # Auto-save to GitHub when configured
        if github_configured():
            with st.spinner("Saving updated inventory to GitHub…"):
                msg = save_master_excel(st.session_state.inventory)
            st.session_state.github_status = msg
            st.info(msg)
        else:
            st.caption(
                "Tip: configure GitHub secrets so the inventory is saved on Cloud automatically."
            )

    if st.session_state.last_logs and not st.session_state.inventory.empty:
        with st.expander("Last processing log"):
            for line in st.session_state.last_logs:
                st.text(line)

# ========================= TAB: Full Inventory ===========================
with tab_inventory:
    st.header("Full Inventory")
    inv = st.session_state.inventory

    if inv.empty:
        st.info("Inventory is empty. Go to **Upload & Add** to load sheets.")
    else:
        # ---- helpers for missing cert detection ----
        def _is_blank_cert(val) -> bool:
            if val is None:
                return True
            try:
                if pd.isna(val):
                    return True
            except (TypeError, ValueError):
                pass
            s = str(val).strip().lower()
            return s in ("", "nan", "none", "null", "<na>", "nat")

        inv = inv.copy()
        inv["_row_id"] = range(len(inv))  # stable id for editing merges
        missing_mask = inv["certificate_no"].map(_is_blank_cert)
        n_missing = int(missing_mask.sum())

        m1, m2, m3 = st.columns(3)
        m1.metric("Total stones", len(inv))
        m2.metric("Missing Certificate #", n_missing)

        if n_missing:
            st.warning(
                f"**{n_missing} stone(s) have no Certificate #.** "
                "Those rows are highlighted below — enter the certificate number in the "
                "**Certificate #** column, then click **Apply certificate edits**."
            )

        # ---- Search + filters (every column) ----
        st.subheader("Search & filters")
        search = st.text_input(
            "🔍 Search all columns (keywords, certificate, stock, shape, links…)",
            placeholder="e.g. oval SI1  or  1535750439  or  GIA",
            key="inv_search",
        )

        with st.expander("Column filters", expanded=False):
            fcols = st.columns(4)
            shapes = sorted([str(s) for s in inv["shape"].dropna().unique()])
            colors = sorted([str(c) for c in inv["color"].dropna().unique()])
            clarities = sorted([str(c) for c in inv["clarity"].dropna().unique()])
            sources = sorted([str(s) for s in inv["source_file"].dropna().unique()])
            labs = sorted([str(s) for s in inv["lab"].dropna().unique()]) if "lab" in inv.columns else []
            polishes = sorted([str(s) for s in inv["polish"].dropna().unique()]) if "polish" in inv.columns else []

            sel_shape = fcols[0].multiselect("Shape", shapes, default=[])
            sel_color = fcols[1].multiselect("Color", colors, default=[])
            sel_clarity = fcols[2].multiselect("Clarity", clarities, default=[])
            sel_source = fcols[3].multiselect("Source file", sources, default=[])

            fcols2 = st.columns(4)
            sel_lab = fcols2[0].multiselect("Lab", labs, default=[]) if labs else []
            sel_polish = fcols2[1].multiselect("Polish", polishes, default=[]) if polishes else []
            only_missing = fcols2[2].checkbox("Only missing Certificate #", value=False)
            only_video = fcols2[3].checkbox("Only with video link", value=False)

            wmin = float(inv["weight"].min()) if inv["weight"].notna().any() else 0.0
            wmax = float(inv["weight"].max()) if inv["weight"].notna().any() else 10.0
            if wmin == wmax:
                wmax = wmin + 0.01
            weight_range = st.slider("Weight (ct)", wmin, wmax, (wmin, wmax))

        # Apply filters
        filtered = inv.copy()
        if search and search.strip():
            tokens = [t.strip().lower() for t in search.split() if t.strip()]
            # every token must appear somewhere in the row (AND across tokens)
            def row_match(row) -> bool:
                blob = " ".join("" if (v is None or (isinstance(v, float) and pd.isna(v))) else str(v) for v in row.values).lower()
                return all(tok in blob for tok in tokens)

            filtered = filtered[filtered.apply(row_match, axis=1)]

        if sel_shape:
            filtered = filtered[filtered["shape"].astype(str).isin(sel_shape)]
        if sel_color:
            filtered = filtered[filtered["color"].astype(str).isin(sel_color)]
        if sel_clarity:
            filtered = filtered[filtered["clarity"].astype(str).isin(sel_clarity)]
        if sel_source:
            filtered = filtered[filtered["source_file"].astype(str).isin(sel_source)]
        if sel_lab:
            filtered = filtered[filtered["lab"].astype(str).isin(sel_lab)]
        if sel_polish:
            filtered = filtered[filtered["polish"].astype(str).isin(sel_polish)]
        if only_missing:
            filtered = filtered[filtered["certificate_no"].map(_is_blank_cert)]
        if only_video:
            filtered = filtered[filtered["video_link"].map(lambda v: is_url(v))]
        filtered = filtered[
            (filtered["weight"].fillna(0) >= weight_range[0])
            & (filtered["weight"].fillna(0) <= weight_range[1])
        ]

        m3.metric("Shown after filters", len(filtered))
        st.caption(f"Showing **{len(filtered)}** of **{len(inv)}** stones")

        # ---- Editable table (certificate_no always editable) ----
        display_cols = [c for c in STANDARD_COLUMNS if c in filtered.columns]
        edit_df = filtered[display_cols + ["_row_id"]].copy()

        def _as_str(v):
            cleaned = clean_value(v)
            return "" if cleaned is None else cleaned

        # Coerce types so data_editor / Arrow never sees mixed object columns
        numeric_cols = {"weight", "price_per_ct", "amount"}
        for col in list(edit_df.columns):
            if col == "_row_id":
                edit_df[col] = pd.to_numeric(edit_df[col], errors="coerce").astype("int64")
            elif col in numeric_cols:
                edit_df[col] = pd.to_numeric(edit_df[col], errors="coerce")
            else:
                edit_df[col] = edit_df[col].map(_as_str).astype("string")

        # Red mark for missing certificate numbers
        miss_flags = edit_df["certificate_no"].map(
            lambda v: "🔴" if _is_blank_cert(v) else ""
        )
        edit_df.insert(0, "Missing Cert", miss_flags.astype("string"))

        column_config = {
            "Missing Cert": st.column_config.TextColumn(
                "🔴",
                disabled=True,
                width="small",
                help="Red mark = Certificate # is missing — please enter it",
            ),
            "certificate_no": st.column_config.TextColumn(
                "Certificate #",
                help="Mandatory — enter GIA/IGI report number",
                width="medium",
            ),
            "stock_no": st.column_config.TextColumn("Stock #", width="medium"),
            "weight": st.column_config.NumberColumn("Weight", format="%.2f"),
            "price_per_ct": st.column_config.NumberColumn("Price/ct", format="%.2f"),
            "amount": st.column_config.NumberColumn("Amount", format="%.2f"),
            "video_link": st.column_config.TextColumn("Video link", width="medium"),
            "image_link": st.column_config.TextColumn("Image link", width="medium"),
            "certificate_link": st.column_config.TextColumn("Cert link", width="medium"),
            "_row_id": st.column_config.NumberColumn("_row_id", disabled=True),
        }

        col_order = ["Missing Cert"] + [c for c in display_cols]

        st.markdown(
            "Edit **Certificate #** (and other fields) directly in the grid. "
            "Then click **Apply edits** to save into the inventory."
        )
        edited = st.data_editor(
            edit_df,
            use_container_width=True,
            height=480,
            hide_index=True,
            num_rows="fixed",
            column_config=column_config,
            column_order=col_order,
            disabled=[c for c in edit_df.columns if c not in (
                "certificate_no", "stock_no", "lab", "shape", "color", "clarity",
                "polish", "symmetry", "fluorescence", "cut", "notes",
                "video_link", "image_link", "certificate_link", "mm_size",
            )],
            key="inventory_editor",
        )

        c_apply, c_save = st.columns(2)
        with c_apply:
            if st.button("✅ Apply edits to inventory", type="primary", use_container_width=True):
                # Merge edited rows back into full inventory by _row_id
                base = st.session_state.inventory.copy()
                base["_row_id"] = range(len(base))
                editable_cols = [
                    "certificate_no", "stock_no", "lab", "shape", "color", "clarity",
                    "polish", "symmetry", "fluorescence", "cut", "notes",
                    "video_link", "image_link", "certificate_link", "mm_size",
                ]
                for _, row in edited.iterrows():
                    rid = row["_row_id"]
                    idx = base.index[base["_row_id"] == rid]
                    if len(idx) == 0:
                        continue
                    i = idx[0]
                    for col in editable_cols:
                        if col in row.index and col in base.columns:
                            val = row[col]
                            # normalize blank cert to None
                            if col in ("certificate_no", "stock_no"):
                                val = clean_value(val)
                            base.at[i, col] = val
                base = base.drop(columns=["_row_id"], errors="ignore")
                st.session_state.inventory = _normalize_df(base)
                still = st.session_state.inventory["certificate_no"].map(_is_blank_cert).sum()
                if still:
                    st.warning(f"Edits applied. **{int(still)}** stone(s) still missing Certificate #.")
                else:
                    st.success("Edits applied. All stones have a Certificate #.")
                st.rerun()

        with c_save:
            if github_configured():
                if st.button("💾 Apply edits & save to GitHub", use_container_width=True):
                    base = st.session_state.inventory.copy()
                    base["_row_id"] = range(len(base))
                    editable_cols = [
                        "certificate_no", "stock_no", "lab", "shape", "color", "clarity",
                        "polish", "symmetry", "fluorescence", "cut", "notes",
                        "video_link", "image_link", "certificate_link", "mm_size",
                    ]
                    for _, row in edited.iterrows():
                        rid = row["_row_id"]
                        idx = base.index[base["_row_id"] == rid]
                        if len(idx) == 0:
                            continue
                        i = idx[0]
                        for col in editable_cols:
                            if col in row.index and col in base.columns:
                                val = row[col]
                                if col in ("certificate_no", "stock_no"):
                                    val = clean_value(val)
                                base.at[i, col] = val
                    base = base.drop(columns=["_row_id"], errors="ignore")
                    st.session_state.inventory = _normalize_df(base)
                    msg = save_master_excel(st.session_state.inventory)
                    st.session_state.github_status = msg
                    st.info(msg)
                    st.rerun()

        st.markdown("#### Download")
        d1, d2, d3 = st.columns(3)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        # strip helper cols for download
        inv_dl = st.session_state.inventory
        filt_dl = filtered.drop(columns=["_row_id"], errors="ignore")
        with d1:
            st.download_button(
                "📥 Master Excel (all)",
                data=_df_to_excel_bytes(inv_dl),
                file_name=f"master_diamonds_{ts}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )
        with d2:
            st.download_button(
                "📥 Master CSV (all)",
                data=_df_to_csv_bytes(inv_dl),
                file_name=f"master_diamonds_{ts}.csv",
                mime="text/csv",
                use_container_width=True,
            )
        with d3:
            st.download_button(
                "📥 Filtered Excel",
                data=_df_to_excel_bytes(filt_dl[display_cols] if not filt_dl.empty else filt_dl),
                file_name=f"filtered_diamonds_{ts}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                disabled=filtered.empty,
            )

# ========================= TAB: Summary ==================================
with tab_summary:
    st.header("Summary Dashboard")
    inv = st.session_state.inventory

    if inv.empty:
        st.info("Inventory is empty. Go to **Upload & Add** to load sheets.")
    else:
        k1, k2, k3, k4, k5 = st.columns(5)
        k1.metric("Total stones", len(inv))
        k2.metric("Total carats", f"{inv['weight'].sum():.2f}")
        avg_price = inv["price_per_ct"].mean()
        k3.metric("Avg $/ct", f"{avg_price:,.0f}" if pd.notna(avg_price) else "—")
        total_val = inv["amount"].sum()
        if not total_val or pd.isna(total_val) or total_val == 0:
            total_val = (inv["weight"] * inv["price_per_ct"]).sum()
        k4.metric("Est. total value", f"${total_val:,.0f}" if pd.notna(total_val) else "—")
        with_video = inv["video_link"].notna().sum()
        k5.metric("With video link", int(with_video))

        st.markdown("---")

        c1, c2, c3 = st.columns(3)
        with c1:
            st.subheader("By Shape")
            shape_ct = (
                inv.groupby("shape", dropna=False)
                .agg(stones=("weight", "count"), carats=("weight", "sum"))
                .reset_index()
                .sort_values("stones", ascending=False)
            )
            st.dataframe(shape_ct, use_container_width=True, hide_index=True)
            st.bar_chart(shape_ct.set_index("shape")["stones"])

        with c2:
            st.subheader("By Color")
            color_ct = (
                inv.groupby("color", dropna=False)
                .agg(stones=("weight", "count"), carats=("weight", "sum"))
                .reset_index()
                .sort_values("stones", ascending=False)
            )
            st.dataframe(color_ct, use_container_width=True, hide_index=True)
            st.bar_chart(color_ct.set_index("color")["stones"])

        with c3:
            st.subheader("By Clarity")
            clar_ct = (
                inv.groupby("clarity", dropna=False)
                .agg(stones=("weight", "count"), carats=("weight", "sum"))
                .reset_index()
                .sort_values("stones", ascending=False)
            )
            st.dataframe(clar_ct, use_container_width=True, hide_index=True)
            st.bar_chart(clar_ct.set_index("clarity")["stones"])

        st.markdown("---")
        st.subheader("Shape × Color pivot (stone count)")
        try:
            pivot = pd.crosstab(
                inv["shape"].fillna("(blank)"),
                inv["color"].fillna("(blank)"),
                margins=True,
                margins_name="Total",
            )
            st.dataframe(pivot, use_container_width=True)
        except Exception:
            st.write("Not enough data for pivot.")

        st.subheader("By source file")
        src = (
            inv.groupby("source_file", dropna=False)
            .agg(stones=("weight", "count"), carats=("weight", "sum"))
            .reset_index()
            .sort_values("stones", ascending=False)
        )
        st.dataframe(src, use_container_width=True, hide_index=True)

        st.markdown("#### Download summary workbook")
        summary_buf = io.BytesIO()
        with pd.ExcelWriter(summary_buf, engine="openpyxl") as writer:
            inv.to_excel(writer, sheet_name="Full Inventory", index=False)
            shape_ct.to_excel(writer, sheet_name="By Shape", index=False)
            color_ct.to_excel(writer, sheet_name="By Color", index=False)
            clar_ct.to_excel(writer, sheet_name="By Clarity", index=False)
            src.to_excel(writer, sheet_name="By Source", index=False)
            try:
                pivot.to_excel(writer, sheet_name="Shape x Color")
            except Exception:
                pass
        st.download_button(
            "📥 Download Summary Workbook (multi-sheet Excel)",
            data=summary_buf.getvalue(),
            file_name=f"diamond_summary_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
